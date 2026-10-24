"""The VM report (instances.views.vm_report): every VM the user may see, with
its addresses and OS, as an Excel workbook."""

from io import BytesIO
from unittest.mock import patch

from accounts.models import UserInstance
from computes.models import Compute
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.contrib.messages import get_messages
from django.test import TestCase
from django.urls import reverse
from instances import addresses
from instances.models import Instance
from lxml import html
from openpyxl import load_workbook

UUID_A = "11111111-2222-3333-4444-00000000000a"
UUID_B = "11111111-2222-3333-4444-00000000000b"
UNKNOWN = "11111111-2222-3333-4444-00000000000c"
COUNTS = {"checked": 2, "no_agent": 1, "unread": 0, "hosts_down": ["host-three"], "hosts_unfinished": []}


def vm(uuid, name, compute_id, state=1, addresses_=(), os="AlmaLinux 10.0", family="Linux"):
    return {"compute_id": compute_id, "uuid": uuid, "name": name, "state": state, "vcpus": 2, "memory": 2048,
            "macs": ["52:54:00:00:00:01"], "os": os, "os_family": family, "addresses": list(addresses_)}


class VmReportTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_superuser("vr-admin", "vr@example.com", "pw")
        self.owner = User.objects.create_user("vr-owner", password="pw")
        self.c1 = Compute.objects.create(name="host-one", hostname="192.0.2.1", login="u", password="p", type=1)
        self.c2 = Compute.objects.create(name="host-two", hostname="192.0.2.2", login="u", password="p", type=1)
        self.a = Instance.objects.create(compute=self.c1, name="vm-a", uuid=UUID_A)
        Instance.objects.create(compute=self.c2, name="vm-b", uuid=UUID_B)
        UserInstance.objects.create(user=self.owner, instance=self.a)

    def report(self, user, vms=(), side_effect=None):
        self.client.force_login(user)
        with patch("instances.addresses.inventory", return_value={"vms": list(vms), **COUNTS},
                   side_effect=side_effect) as inventory:
            response = self.client.get(reverse("instances:vm_report"))
        return response, inventory

    def sheets(self, response):
        book = load_workbook(BytesIO(response.content))
        return ([list(row) for row in book["VMs"].iter_rows(values_only=True)],
                [row[0] for row in book["Coverage"].iter_rows(values_only=True)])

    def test_an_administrator_gets_every_vm_with_hosts_and_owners(self):
        response, inventory = self.report(self.admin, [
            vm(UUID_A, "vm-a", self.c1.id, addresses_=[("192.0.2.10", ["ARP", "agent"]), ("2001:db8::10", ["agent"]),
                                                       ("fe80::1", ["agent"])]),  # link-local: not an address to report
            vm(UNKNOWN, "made-with-virsh", self.c2.id, state=5, os="", family=""),
        ])
        self.assertEqual({(t.compute.id, t.uuids) for t in inventory.call_args.args[0]},
                         {(self.c1.id, None), (self.c2.id, None)})
        self.assertEqual(response["Content-Type"],
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertRegex(response["Content-Disposition"], r'attachment; filename="vms-\d{8}-\d{4}\.xlsx"')
        rows, coverage = self.sheets(response)
        self.assertEqual(rows[0], ["VM", "State", "Host", "OS", "OS family", "vCPUs", "Memory (MiB)", "IP addresses",
                                   "Address sources", "MAC addresses", "Owners", "In the panel"])
        self.assertEqual(rows[1], ["vm-a", "Active", "host-one", "AlmaLinux 10.0", "Linux", 2, 2048,
                                   "192.0.2.10, 2001:db8::10", "ARP, agent", "52:54:00:00:00:01", "vr-owner", "yes"])
        self.assertEqual(rows[2][:3] + rows[2][-1:], ["made-with-virsh", "Off", "host-two", "no"])
        self.assertIn("Hosts not reached: host-three", coverage)

    def test_only_superusers_get_it(self):
        viewer = get_user_model().objects.create_user("vr-viewer", password="pw")
        viewer.user_permissions.add(Permission.objects.get(codename="view_instances"))
        for user in (self.owner, viewer):
            with self.subTest(user=user.username):
                response, inventory = self.report(user)
                self.assertEqual(response.status_code, 403)
                inventory.assert_not_called()

    def test_the_administration_menu_offers_it_to_superusers_only(self):
        for user, offered in ((self.admin, True), (self.owner, False)):
            with self.subTest(user=user.username):
                self.client.force_login(user)
                page = html.fromstring(self.client.get(reverse("accounts:profile")).content)
                links = page.xpath("//div[@aria-labelledby='administration']//a/@href")
                self.assertEqual(reverse("instances:vm_report") in links, offered)

    def test_too_many_searches_return_to_the_list(self):
        response, _ = self.report(self.admin, side_effect=addresses.Busy)
        self.assertRedirects(response, reverse("instances:index"), fetch_redirect_response=False)
        self.assertIn("try again shortly", " ".join(str(m) for m in get_messages(response.wsgi_request)))

    def test_text_that_looks_like_a_formula_stays_text(self):
        # names come from VM owners and guests: Excel must not run them
        response, _ = self.report(self.admin, [vm(UUID_A, '=HYPERLINK("http://example.com","x")', self.c1.id,
                                                  os="+cmd|' /C calc'!A0")])
        cells = load_workbook(BytesIO(response.content))["VMs"]["A2":"D2"][0]
        self.assertEqual([c.data_type for c in cells], ["s", "s", "s", "s"])
        self.assertEqual(cells[0].value, '=HYPERLINK("http://example.com","x")')
