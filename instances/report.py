"""An Excel workbook: rows under a bold, filterable header, and a sheet of notes."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

MAX_WIDTH = 60


def workbook(title, header, rows, notes_title, notes):
    """The .xlsx file's bytes. Every text cell stays text: openpyxl would
    write one starting with "=" as a formula, and the text comes from VM
    names and guests."""
    book = Workbook()
    sheet = book.active
    sheet.title = title
    for values in [header, *rows]:
        # a control character (a name typed at a failed login) cannot be in the file: written as \x01
        sheet.append([ILLEGAL_CHARACTERS_RE.sub(lambda m: f"\\x{ord(m.group()):02x}", value)
                      if isinstance(value, str) else value for value in values])
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(sheet.columns, start=1):
        width = max(len(str(cell.value or "")) for cell in column) + 2
        sheet.column_dimensions[get_column_letter(index)].width = min(width, MAX_WIDTH)
    notes_sheet = book.create_sheet(notes_title)
    for line in notes:
        notes_sheet.append([line])
    notes_sheet.column_dimensions["A"].width = MAX_WIDTH * 2
    out = BytesIO()
    book.save(out)
    return out.getvalue()
