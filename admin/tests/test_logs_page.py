"""The Logs page searches every row on the server, can show the failed ones
only, says the year and time zone, links a VM's rows to the VM, and downloads
what it shows as an Excel workbook."""

from io import BytesIO
from unittest.mock import patch

from computes.models import Compute
from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse
from instances.models import Instance
from logs.models import Logs
from lxml import html
from openpyxl import load_workbook

UUID = "11111111-2222-3333-4444-555555555555"


@override_settings(TIME_ZONE="UTC")
class LogsPageTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_superuser("lg-admin", "lg@example.com", "pw")
        self.client.force_login(self.admin)
        Logs.objects.all().delete()
        # the row to find is the oldest: past the first page
        Logs.objects.create(user="alice", host="host1", instance="web", instance_uuid=UUID,
                            message="VM web: delete", success=False, ip="192.0.2.7")
        Logs.objects.bulk_create(Logs(user="bob", instance="db", message=f"VM db: power on {n}", success=True)
                                 for n in range(150))
        compute = Compute.objects.create(name="host1", hostname="192.0.2.1", login="u", password="p", type=1)
        self.vm = Instance.objects.create(compute=compute, name="web", uuid=UUID)

    def page(self, **query):
        return html.fromstring(self.client.get(reverse("admin:logs"), query).content)

    def rows(self, page):
        return [[" ".join(td.text_content().split()) for td in tr.xpath("td")] for tr in page.xpath("//tbody/tr")]

    def test_a_search_finds_a_row_on_any_page(self):
        for query in ("alice", "WEB: DELETE", "192.0.2.7", "host1"):
            with self.subTest(query=query):
                rows = self.rows(self.page(q=query))
                self.assertEqual([row[2] for row in rows], ["alice"])

    def test_a_search_is_text_not_a_pattern(self):
        Logs.objects.create(user="carol", instance="", message="disk (vda) [full]", success=True)
        self.assertEqual([row[2] for row in self.rows(self.page(q="(vda) ["))], ["carol"])

    def test_the_failed_rows_only_and_marked(self):
        page = self.page(failed="1")
        self.assertEqual([row[2] for row in self.rows(page)], ["alice"])
        self.assertIn("table-danger", page.xpath("//tbody/tr")[0].get("class"))

    def test_the_pages_keep_the_search(self):
        page = self.page(q="bob")
        self.assertTrue(page.xpath("//a[contains(@href, 'page=2') and contains(@href, 'q=bob')]"))

    def test_dates_have_the_year_and_the_time_zone(self):
        page = self.page(q="alice")
        self.assertIn("Date (UTC)", [" ".join(th.text_content().split()) for th in page.xpath("//thead//th")])
        log = Logs.objects.get(user="alice")
        self.assertEqual(self.rows(page)[0][1], log.date.strftime("%Y-%m-%d %H:%M:%S"))

    def test_a_vm_row_links_the_vm_while_it_is_there(self):
        page = self.page(q="alice")
        self.assertEqual(page.xpath("//tbody/tr/td/a/@href"), [reverse("instances:instance", args=[self.vm.pk])])
        self.vm.delete()
        self.assertEqual(self.page(q="alice").xpath("//tbody/tr/td/a"), [])

    def test_the_rows_shown_download_as_a_workbook(self):
        Logs.objects.create(user="mallory", instance="", message="=HYPERLINK(1)", success=False)
        response = self.client.get(reverse("admin:logs_export"), {"failed": "1"})
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("attachment", response["Content-Disposition"])
        sheet = load_workbook(BytesIO(response.content)).active
        values = list(sheet.values)
        self.assertEqual(values[0], ("Date (UTC)", "User", "Host", "VM", "VM UUID", "Message", "Result", "IP"))
        self.assertEqual([row[1] for row in values[1:]], ["mallory", "alice"])
        # a message starting with = stays text
        self.assertEqual(sheet.cell(row=2, column=6).data_type, "s")
        page = self.page(q="alice", failed="1")
        self.assertEqual(page.xpath("//a[contains(@href, 'export')]/@href"),
                         [reverse("admin:logs_export") + "?q=alice&failed=1"])

    def test_only_an_administrator_downloads(self):
        self.client.force_login(get_user_model().objects.create_user("lg-user", password="pw"))
        self.assertEqual(self.client.get(reverse("admin:logs_export")).status_code, 403)

    def test_a_control_character_in_a_row_is_shown_not_a_failure(self):
        # a failed login writes the name it was tried with, as typed
        Logs.objects.create(user="eve\x01", instance="", message="Login", success=False)
        response = self.client.get(reverse("admin:logs_export"), {"q": "eve"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(load_workbook(BytesIO(response.content)).active.cell(row=2, column=2).value, "eve\\x01")

    def test_the_download_stops_at_its_row_limit_and_says_so(self):
        with patch("admin.views.EXPORT_ROWS", 3):
            book = load_workbook(BytesIO(self.client.get(reverse("admin:logs_export")).content))
        self.assertEqual(book.active.max_row, 4)  # the header and three rows
        notes = [row[0] for row in book["About"].values]
        self.assertIn("Only the newest 3 rows: narrow the search for the others", notes)

